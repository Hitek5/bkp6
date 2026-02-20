from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def create_template(path: str = "templates/input_template.xlsx") -> Path:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    ws_packs = wb.active
    ws_packs.title = "Комплекты"
    ws_packs.append([
        "DocID",
        "KKS",
        "Наименование",
        "Спец",
        "Здание",
        "Система",
        "Отметка",
        "Статус",
        "Типовой",
    ])
    ws_packs.append(["DOC-001", "10KZ01", "Пример комплекта", "KZ", "BLD-01", "SYS-01", "+0.000", "ACTIVE", "N"])

    ws_matrix = wb.create_sheet("Матрица_триггеров")
    ws_matrix.append(["Триггер", "KZ", "AR", "OV", "VK"])
    ws_matrix.append(["SEIS_7_9", 1, 1, 0, 0])

    ws_rules = wb.create_sheet("Правила")
    ws_rules.append(
        [
            "RuleID",
            "Триггер",
            "Спец",
            "Здания",
            "Системы",
            "Отметки",
            "Режим(INCLUDE/EXCLUDE)",
            "Комментарий",
            "Автор",
            "Дата",
        ]
    )
    ws_rules.append(["R-001", "SEIS_7_9", "KZ", "BLD-*", "SYS-*", "*", "INCLUDE", "Базовое включение", "system", "2026-01-01"])

    ws_result = wb.create_sheet("Результат")
    ws_result.append(["DocID", "KKS", "Наименование", "Спец", "Здание", "Система", "Отметка", "Причина_попадания"])

    wb.save(out)
    return out


if __name__ == "__main__":
    output = create_template()
    print(f"Template created: {output}")
